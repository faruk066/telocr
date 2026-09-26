"""JSON satirlari -> dogrulamali, bicimlendirilmis .xlsx uretimi."""
from __future__ import annotations

import logging
import os
import re
import time
from collections import Counter

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

logger = logging.getLogger(__name__)

COLUMNS = [
    "BİNA",
    "DAİRE",
    "MARKA",
    "SAYAÇ TÜRÜ",
    "DURUM",
    "TUTAR (TL)",
    "ESKİ ENDEKS",
    "YENİ SERİ NO (BARKOD)",
    "NOTLAR",
    "ENLEM",
    "BOYLAM",
    "KONTROL NOTU",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
ORANGE_FILL = PatternFill("solid", fgColor="FCE4D6")
GREEN_FONT = Font(color="006100")
RED_FONT = Font(color="9C0006")
ORANGE_FONT = Font(color="9C4221")

# Mukerrer seri gruplari icin palet: ayni seri -> ayni renk.
# Acik pastel zemin + koyu yazi (okunabilirlik icin).
DUP_PALETTE: list[tuple[str, str]] = [
    ("FFD9E1", "9C0A2E"),  # pembe
    ("D9E8FF", "1A3FAA"),  # mavi
    ("FFE8C2", "9A5B00"),  # turuncu
    ("D9F2DF", "0E6B2E"),  # yesil
    ("E8D9FF", "5B21B6"),  # mor
    ("FFF3B0", "7A6200"),  # sari
    ("D6F0F0", "0B6E6E"),  # turkuaz
    ("F2D9D9", "8C1D1D"),  # bordo
    ("E2E8F0", "334155"),  # gri-mavi
    ("DCFCE7", "166534"),  # acik yesil
    ("FEF3C7", "92400E"),  # amber
    ("E0E7FF", "3730A3"),  # indigo
]

WARN_TEXT = "⚠️ Kontrol Edilmeli"


def _digits(s: object) -> str:
    return re.sub(r"\D", "", str(s or ""))


def _looks_like_17_swap(a: str, b: str) -> bool:
    """Iki seri no arasindaki TUM farklar 1<->7 cifti ve EN AZ 2 hanedeyse True.

    Turkce/Avrupa yaziminda cizgili '7' cogu zaman '1' sanilir; boylece
    30507777 -> 30501111 gibi, birden fazla hanede olabilen sapmalar olusur ve
    ayni seri numarasi birden fazla daireye yazilir. Tek hanelik 1/7 farki
    (ornegin 30501777 / 30501771) rastlantusal olabilecegi icin esik de 2'dir.
    """
    if len(a) != len(b) or a == b:
        return False
    diffs = [(x, y) for x, y in zip(a, b) if x != y]
    if len(diffs) < 2:
        return False
    return all({x, y} == {"1", "7"} for x, y in diffs)


# --- Seri no alanina karisan sahte numaralar (telefon / T.C. kimlik) ----------
# Sayac barkodu HER ZAMAN tam 8 hanedir ('30...' ultrasonik, '23'/'24' mekanik,
# '80...' sicak su). Formlara musterinin telefonu ya da kimlik numarasi
# karistiginda bunlar barkod olarak KAYDEDILMEZ; NOTLAR'a tasinir.

def _valid_turkish_id(n: str) -> bool:
    """T.C. kimlik numarasi checksum kontrolu (11 hane, algoritmik olarak tutarli)."""
    if len(n) != 11 or not n.isdigit() or n[0] == "0":
        return False
    d = [int(c) for c in n]
    if ((d[0] + d[2] + d[4] + d[6] + d[8]) * 7 - (d[1] + d[3] + d[5] + d[7])) % 10 != d[9]:
        return False
    return sum(d[:10]) % 10 == d[10]


def _phone_core(digits: str) -> str:
    """Ulke/sabit-hat oneklerini soyup 10 haneli yerel numaraya indirger.

    '00905321234567' -> '5321234567', '905321234567' -> '5321234567',
    '05321234567' -> '5321234567', '5321234567' -> '5321234567'.
    """
    s = digits
    while s.startswith("00"):
        s = s[2:]                              # 0090... -> 90...
    if s.startswith("90") and len(s) > 10:
        s = s[2:]                              # 90532... -> 532...
    if s.startswith("0") and len(s) == 11:
        s = s[1:]                              # 0532... -> 532...
    return s


def _bad_serial_reason(digits: str) -> str:
    """Bu rakam dizisi barkod OLAMAZSA nedenini doner; gecerliyse bos string.

    Kurallar: 8 hane -> her zaman kabul. 5-7 hane -> kabul ama uyari
    (kesilmis/eksik okunmus barkod olabilir). 9+ hane -> barkod DEGILDIR;
    telefon/kimlik deselerine gore tiplendirilir.
    """
    n = len(digits)
    if n <= 8 or not digits.isdigit():
        return ""
    # 11 hane telefonla birebir ayni uzunlukta; kimlik kontrolu onekleri
    # soyulmadan once yapilir (gecerli T.C. no '0' ile baslamaz).
    if n == 11 and _valid_turkish_id(digits):
        return "T.C. kimlik numarasi"
    core = _phone_core(digits)
    # 10 haneli yerel numara: 5xx mobil, 2xx/3xx/4xx sabit hat, 850/444 hizmet.
    if len(core) == 10 and core[0] in "23458":
        return "telefon numarasi"
    return f"{n} haneli numara (barkod 8 hane olmali)"


def _format_phone(digits: str) -> str:
    """Rakam yiginini okunur telefona cevir: 5321234567 -> 0532 123 45 67."""
    body = _phone_core(digits)
    if len(body) != 10:
        return digits
    return f"0{body[:3]} {body[3:6]} {body[6:8]} {body[8:]}"


# KONTROL NOTU icinde bu isaret bulunan satirlar Excel'de ayrica vurgulanir.
SERIAL_INVALID_MARK = "SERI NO DEGIL"


def sanitize_serials(rows: list[dict]) -> tuple[list[dict], int]:
    """Seri no alanindaki telefon/kimlik numaralarini cikarip NOTLAR'a tasir.

    Dondurur: (temizlenmis satirlar, cikarilan sahte seri sayisi). Gercek barkod
    8 hane oldugu icin 9+ haneli deger 'seri var -> Tamamlandi' kuralini da
    yaniltir; bu yuzden alan bosaltilir, deger NOTLAR'a yazilir, DURUM
    'Kontrol Edilmeli' yapilir ve satira 'SERI NO DEGIL' isareti konur.
    """
    out: list[dict] = []
    removed = 0
    for r in rows:
        row = dict(r)
        digits = _digits(row.get("YENİ SERİ NO (BARKOD)", ""))
        reason = _bad_serial_reason(digits)
        if not reason:
            out.append(row)
            continue
        shown = _format_phone(digits) if reason == "telefon numarasi" else digits
        note = str(row.get("NOTLAR", "") or "").strip()
        addition = f"Seri no alanindaki {shown} {reason}"
        row["NOTLAR"] = f"{note} / {addition}" if note else addition
        row["YENİ SERİ NO (BARKOD)"] = ""
        row["DURUM"] = "Kontrol Edilmeli"
        row["_seri_no_degil"] = shown
        removed += 1
        logger.warning("Seri no alanindaki %s cikarildi (%s).", shown, reason)
        out.append(row)
    return out, removed


def validate_rows(rows: list[dict]) -> list[dict]:
    """8-hane + mukerrer + 1/7 karismasi + sahte seri (telefon/kimlik) kontrolu.

    Sorunlu satira KONTROL NOTU ekler, islemi durdurmaz. Mukerrer seri no,
    yalnizca 1/7 ile ayrilan seriler ve seri no sanilan telefon numaralari
    isaretlenir (cizgili '7' yazan eller '1' okunabiliyor).
    """
    serials = [_digits(r.get("YENİ SERİ NO (BARKOD)", "")) for r in rows]
    counts = Counter(s for s in serials if s)
    daire_counts = Counter(_daire_key(r.get("DAİRE", "")) for r in rows)
    daire_by_serial: dict[str, list[str]] = {}
    for r, s in zip(rows, serials):
        if s:
            daire_by_serial.setdefault(s, []).append(str(r.get("DAİRE", "")).strip())
    # Tum hanelerinde 1<->7 farki olan seri ciftleri
    near_miss: dict[str, set[str]] = {}
    uniq = sorted(daire_by_serial)
    for i, a in enumerate(uniq):
        for b in uniq[i + 1:]:
            if _looks_like_17_swap(a, b):
                near_miss.setdefault(a, set()).add(b)
                near_miss.setdefault(b, set()).add(a)
    out: list[dict] = []
    for r, serial in zip(rows, serials):
        warnings: list[str] = []
        fake = str(r.get("_seri_no_degil", "") or "")
        if fake:
            warnings.append(f"{SERIAL_INVALID_MARK}: {fake} barkod degil, NOTLAR'a tasindi")
        if serial:
            if len(serial) != 8:
                warnings.append(f"Seri no 8 haneli degil ({serial})")
            if counts[serial] > 1:
                ds = ", ".join(daire_by_serial[serial])
                warnings.append(
                    f"Mukerrer seri no ({serial}) - daire {ds}; 1/7 rakam karismasi olabilir")
            elif serial in near_miss:
                others = ", ".join(sorted(near_miss[serial]))
                warnings.append(
                    f"Seri {serial} ile {others} yalnizca 1/7 farki iceriyor; "
                    "cizgili 7 / dikey 1 ayrimini kontrol et")
        if daire_counts[_daire_key(r.get("DAİRE", ""))] > 1:
            warnings.append(f"Ayni daire birden fazla satirda ({r.get('DAİRE', '')})")
        # Mevcut NOTLAR'i koru, kontrol notunu ayri sutuna yaz
        row = {c: r.get(c, "") for c in COLUMNS if c != "KONTROL NOTU"}
        row["KONTROL NOTU"] = "; ".join(warnings) if warnings else ""
        if warnings:
            row["KONTROL NOTU"] = f"{WARN_TEXT}: " + row["KONTROL NOTU"]
        out.append(row)
    return out


def _building_prefix(rows: list[dict]) -> str:
    """Dosya adi oneki: en sik gecen BINA degeri, guvenli dosya adiyla.

    Ornek: BINA='Mega Park B Blok' -> 'Mega_Park_B_Blok_'.
    Bina bos ya da hepsi farkliysa bos string doner (eski ad korunur).
    """
    names = [str(r.get("BİNA", "") or "").strip() for r in rows]
    names = [n for n in names if n]
    if not names:
        return ""
    top, count = Counter(names).most_common(1)[0]
    if count < 2 and len(set(names)) > 1:
        return ""  # tutarli bir bina adi yok -> tahmin yurutme
    safe = re.sub(r"[^\w\-]+", "_", top, flags=re.UNICODE).strip("_")
    safe = re.sub(r"_+", "_", safe)
    return f"{safe}_" if safe else ""


def _daire_key(v: object) -> str:
    """Daire karsilastirmasi icin normalize anahtar: 'Daire 7' -> '7'."""
    s = str(v or "").strip().lower()
    m = re.search(r"\d+", s)
    return m.group(0) if m else s


def _serial_score(r: dict) -> tuple[int, int]:
    """Ayni dairenin iki kaydi capisirsa hangisi kazanir: (seri_var_mi, seri_uzunlugu)."""
    s = _digits(r.get("YENİ SERİ NO (BARKOD)", ""))
    return (1 if s else 0, len(s))


def make_excel(rows: list[dict], out_dir: str | None = None) -> str:
    """Satirlari .xlsx'e yazar, dosya yolunu doner.

    Model cok fotolu birlestirmeyi atlamissa (ayni daire 2x), kod duzeltir:
    her daireden en iyi kayit tutulur, daire sirasina gore siralanir.
    """
    out_dir = out_dir or config.TEMP_DIR
    os.makedirs(out_dir, exist_ok=True)

    # 1) Seri no alanina karisan telefon/kimlik numaralari BIRAKILMAZ:
    #    birlestirme ONCE temizlenmis satirlar uzerinde calisir (boylece
    #    11 haneli telefon, 8 haneli gercek barkoddan 'daha net' sayilip
    #    kazanmaz) ve sahte degerler NOTLAR'a tasinir.
    sanitized, removed = sanitize_serials(rows)
    if removed:
        logger.info("Seri no alanindan %d sahte deger (telefon/kimlik) cikarildi.", removed)

    merged = merge_duplicate_daires(sanitized)
    if len(merged) != len(sanitized):
        logger.info("Cok-fotolu birlestirme: %d -> %d satir (mukerrer daireler tekillendi).",
                    len(sanitized), len(merged))
    validated = validate_rows(merged)
    df = pd.DataFrame(validated, columns=COLUMNS)

    fname = f"{_building_prefix(validated)}sayac_listesi_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(out_dir, fname)
    df.to_excel(path, index=False, sheet_name="Sayaçlar")

    _format_workbook(path)
    warn_count = sum(1 for r in validated if r.get("KONTROL NOTU"))
    logger.info("Excel olusturuldu: %s (%d satir, %d uyarili).", path, len(validated), warn_count)
    return path


def merge_duplicate_daires(rows: list[dict]) -> list[dict]:
    """Ayni DAIRE'ye ait birden fazla kaydi teke indir + daireye gore sirala.

    Secim: seri numarali (dolu) kayit bosa tercih edilir; ikisi de doluysa
    daha uzun/net seri numarali kazanir. Kaybeden satirin NOTLAR'i doluysa
    ve kazananinki bossa, notlar birlestirilir (bilgi kaybi olmaz).
    """
    best: dict[str, dict] = {}
    order: list[str] = []
    for r in rows:
        key = _daire_key(r.get("DAİRE", ""))
        if not key:
            order.append(f"__noid_{len(order)}")
            best[order[-1]] = dict(r)
            continue
        if key not in best:
            best[key] = dict(r)
            order.append(key)
            continue
        cur = best[key]
        if _serial_score(r) > _serial_score(cur):
            winner, loser = dict(r), cur
        else:
            winner, loser = cur, dict(r)
        # Not kaybi olmasin: kazananin NOTLAR'i bossa, kaybedeninkini ekle
        w_note = str(winner.get("NOTLAR", "") or "").strip()
        l_note = str(loser.get("NOTLAR", "") or "").strip()
        if not w_note and l_note:
            winner["NOTLAR"] = l_note
        elif w_note and l_note and l_note.lower() not in w_note.lower():
            winner["NOTLAR"] = f"{w_note} / {l_note}"
        best[key] = winner
    merged = [best[k] for k in order]

    def _sort_key(r: dict) -> tuple[int, str]:
        m = re.search(r"\d+", str(r.get("DAİRE", "") or ""))
        return (int(m.group(0)) if m else 10**9, str(r.get("DAİRE", "")))

    merged.sort(key=_sort_key)
    return merged


def _format_workbook(path: str) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb["Sayaçlar"]

    # Baslik satiri
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # DURUM kolonu kosullu renklendirme (deger bazli dogrudan boyama - basit ve guvenilir)
    try:
        durum_idx = COLUMNS.index("DURUM") + 1
    except ValueError:
        durum_idx = 5
    try:
        seri_idx = COLUMNS.index("YENİ SERİ NO (BARKOD)") + 1
    except ValueError:
        seri_idx = 8
    try:
        notu_idx = COLUMNS.index("KONTROL NOTU") + 1
    except ValueError:
        notu_idx = 0
    dup_groups = _duplicate_serial_groups(ws, seri_idx)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        cell = row[durum_idx - 1]
        val = str(cell.value or "").strip().lower()
        if val == "tamamlandı":
            cell.fill = GREEN_FILL
            cell.font = GREEN_FONT
        elif val == "iptal":
            cell.fill = RED_FILL
            cell.font = RED_FONT
        elif val == "faruk":
            cell.fill = YELLOW_FILL
        # Seri no alanindaki telefon/kimlik numarasi cikarildiysa turuncu vurgu
        if notu_idx and SERIAL_INVALID_MARK in str(row[notu_idx - 1].value or ""):
            cell.fill = ORANGE_FILL
            cell.font = ORANGE_FONT
    _paint_duplicate_serials(ws, seri_idx, dup_groups)

    # Otomatik sutun genisligi (max 40)
    for i, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        max_len = len(col)
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 4, 40)

    _add_dup_legend(ws)

    wb.save(path)


def _duplicate_serial_groups(ws, seri_idx: int) -> dict[str, list[int]]:
    """Ayni seri noyu paylasan satir numaralarini grupla (sadece 2+ gruplar)."""
    groups: dict[str, list[int]] = {}
    for r in range(2, ws.max_row + 1):
        serial = _digits(ws.cell(r, seri_idx).value)
        if serial:
            groups.setdefault(serial, []).append(r)
    return {s: rows for s, rows in groups.items() if len(rows) > 1}


def _paint_duplicate_serials(ws, seri_idx: int, dup_groups: dict[str, list[int]]) -> None:
    """Her mukerrer seri grubunun SERI + DAIRE hucrelerini ayni renge boya.

    Ayni seriler gozle eslesir; KONTROL NOTU'ndaki yazi da korunur.
    """
    try:
        daire_idx = COLUMNS.index("DAİRE") + 1
    except ValueError:
        daire_idx = 2
    for i, (serial, rows) in enumerate(sorted(dup_groups.items())):
        bg, fg = DUP_PALETTE[i % len(DUP_PALETTE)]
        fill = PatternFill("solid", fgColor=bg)
        font = Font(color=fg, bold=True)
        for r in rows:
            for c in (seri_idx, daire_idx):
                cell = ws.cell(r, c)
                cell.fill = fill
                cell.font = font


def _add_dup_legend(ws) -> None:
    """Tablonun altina mukerrer renk aciklamasi ekle (grup varsa)."""
    try:
        seri_idx = COLUMNS.index("YENİ SERİ NO (BARKOD)") + 1
        daire_idx = COLUMNS.index("DAİRE") + 1
    except ValueError:
        return
    dup_groups = _duplicate_serial_groups(ws, seri_idx)
    if not dup_groups:
        return
    start = ws.max_row + 2
    ws.cell(start, 1, value="MÜKERRER RENK LEJANTI (aynı renk = aynı seri no):").font = Font(bold=True)
    for i, (serial, rows) in enumerate(sorted(dup_groups.items())):
        bg, fg = DUP_PALETTE[i % len(DUP_PALETTE)]
        r = start + 1 + i
        daires = ", ".join(str(ws.cell(rr, daire_idx).value) for rr in rows)
        cell = ws.cell(r, 1, value=f"■ {serial} → Daire: {daires}")
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(color=fg, bold=True)
